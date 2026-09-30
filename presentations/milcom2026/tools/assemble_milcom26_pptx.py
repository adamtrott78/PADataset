#!/usr/bin/env python3
"""Assemble the approved MILCOM 2026 slide renders into the official PPTX template.

The final deck uses full-slide 1920x1080 PNG renders for pixel-faithful visual
reproduction and embeds the canonical spoken script in PowerPoint speaker notes.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "assets" / "template" / "milcom26-ppt-template_v1_25feb26_jb.pptx"
DEFAULT_SCRIPT = ROOT / "SCRIPT.md"
DEFAULT_SPEC = ROOT / "SVG_PRODUCTION_SPEC.md"


def _plain_notes(text: str) -> str:
    """Convert the small amount of Markdown in narration blocks to plain notes text."""
    text = text.strip()
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text, flags=re.S)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    lines = [line.rstrip() for line in text.splitlines()]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    return "\n".join(lines)


def _slide_blocks(markdown: str) -> dict[int, str]:
    matches = list(re.finditer(r"(?m)^# Slide\s+(\d+)\s+—.*$", markdown))
    blocks: dict[int, str] = {}
    for idx, match in enumerate(matches):
        n = int(match.group(1))
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(markdown)
        blocks[n] = markdown[start:end]
    return blocks


def _extract_script_block(block: str) -> str:
    m = re.search(r"(?m)^##\s+(?:Script|Consolidated script)\s*$", block)
    if not m:
        raise ValueError("No Script/Consolidated script heading found")
    remainder = block[m.end():]
    next_heading = re.search(r"(?m)^##\s+", remainder)
    text = remainder[: next_heading.start()] if next_heading else remainder
    return _plain_notes(text)


def _extract_slide1(spec_text: str) -> str:
    m = re.search(r"(?m)^##\s+Opening script\s*$", spec_text)
    if not m:
        raise ValueError("Slide 1 opening script heading not found")
    remainder = spec_text[m.end():]
    end = re.search(r"(?m)^---\s*$", remainder)
    section = remainder[: end.start()] if end else remainder
    quoted: list[str] = []
    for line in section.splitlines():
        if line.startswith(">"):
            quoted.append(line[1:].lstrip())
        elif quoted and not line.strip():
            quoted.append("")
    return _plain_notes("\n".join(quoted))


def load_notes(script_path: Path, spec_path: Path) -> dict[int, str]:
    script_text = script_path.read_text(encoding="utf-8")
    spec_text = spec_path.read_text(encoding="utf-8")

    notes: dict[int, str] = {1: _extract_slide1(spec_text)}
    blocks = _slide_blocks(script_text)
    for n in range(2, 13):
        if n not in blocks:
            raise ValueError(f"SCRIPT.md is missing Slide {n}")
        notes[n] = _extract_script_block(blocks[n])

    missing = [n for n in range(1, 13) if not notes.get(n, "").strip()]
    if missing:
        raise ValueError(f"Empty notes for slides: {missing}")
    return notes


def _remove_all_slide_shapes(slide) -> None:
    for shape in list(slide.shapes):
        el = shape._element
        el.getparent().remove(el)


def _set_notes(slide, text: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        raise RuntimeError("Notes slide does not contain a body placeholder")
    tf.clear()
    tf.text = text


def _validate_pngs(png_dir: Path) -> list[Path]:
    pngs = sorted(png_dir.glob("slide??_*.png"))
    if len(pngs) != 12:
        raise ValueError(f"Expected 12 slide PNGs in {png_dir}, found {len(pngs)}")
    for path in pngs:
        with Image.open(path) as im:
            if im.size != (1920, 1080):
                raise ValueError(f"{path.name}: expected 1920x1080, got {im.size}")
    return pngs


def assemble(template: Path, png_dir: Path, output: Path, script: Path, spec: Path) -> None:
    if not template.exists():
        raise FileNotFoundError(template)
    pngs = _validate_pngs(png_dir)
    notes = load_notes(script, spec)

    prs = Presentation(str(template))
    if len(prs.slides) > 12:
        raise RuntimeError(
            f"Official template unexpectedly has {len(prs.slides)} slides; refusing to truncate"
        )

    ratio = prs.slide_width / prs.slide_height
    if abs(ratio - (16 / 9)) > 0.002:
        raise RuntimeError(f"Official template is not 16:9: ratio={ratio:.6f}")

    # Preserve the official template package/master/layout/theme lineage. Existing
    # template slides are reused; additional slides use the least-populated
    # official layout. The approved slide render is then placed edge-to-edge.
    blankish_layout = min(prs.slide_layouts, key=lambda layout: len(layout.shapes))
    while len(prs.slides) < 12:
        prs.slides.add_slide(blankish_layout)

    for n, (slide, png) in enumerate(zip(prs.slides, pngs), start=1):
        _remove_all_slide_shapes(slide)
        slide.shapes.add_picture(
            str(png),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        _set_notes(slide, notes[n])

    props = prs.core_properties
    props.title = "DQNGuard: Towards Open-World RF Preliminary-Action Detection"
    props.subject = "MILCOM 2026 conference presentation"
    props.author = "Adam Trott"
    props.keywords = "MILCOM 2026, DQNGuard, open-set recognition, RF, Preliminary Actions"
    props.comments = (
        "Assembled from the official MILCOM 2026 template with approved 1920x1080 "
        "slide renders and canonical speaker notes."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))

    # Reopen and verify slide count, dimensions, one full-slide picture per slide,
    # and speaker-note presence after serialization.
    check = Presentation(str(output))
    if len(check.slides) != 12:
        raise RuntimeError(f"Serialized deck has {len(check.slides)} slides, expected 12")
    if check.slide_width != prs.slide_width or check.slide_height != prs.slide_height:
        raise RuntimeError("Serialized deck dimensions changed")

    for n, slide in enumerate(check.slides, start=1):
        pictures = [
            shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        if len(pictures) != 1:
            raise RuntimeError(
                f"Slide {n}: expected one full-slide picture, found {len(pictures)}"
            )
        pic = pictures[0]
        if (pic.left, pic.top, pic.width, pic.height) != (
            0,
            0,
            check.slide_width,
            check.slide_height,
        ):
            raise RuntimeError(f"Slide {n}: picture does not fill slide")
        tf = slide.notes_slide.notes_text_frame
        actual_notes = tf.text.strip() if tf is not None else ""
        if not actual_notes:
            raise RuntimeError(f"Slide {n}: speaker notes missing after save")

    print(f"PASS: assembled {output}")
    print("PASS: 12 slides, 16:9, one exact full-slide render per slide")
    print("PASS: canonical speaker notes embedded on Slides 1–12")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--png-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--script", type=Path, default=DEFAULT_SCRIPT)
    parser.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    args = parser.parse_args()
    assemble(args.template, args.png_dir, args.output, args.script, args.spec)


if __name__ == "__main__":
    main()
